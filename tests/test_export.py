from io import BytesIO

from docx import Document

from core.models import TaskRequest
from services.export import build_docx, build_plain_text


def test_plain_text_export_removes_markdown_heading_markers() -> None:
    request = TaskRequest(scene="非遗问答", raw_request="介绍粤剧")
    answer = "## 一句话认识\n粤剧。\n\n### 核心看点\n- 唱腔"
    text = build_plain_text(request, answer, "### 本次检索资料\n- 项目知识库 · 粤剧")

    assert "# 粤见非遗生成结果" not in text
    assert "## 一句话认识" not in text
    assert "### 核心看点" not in text
    assert "一句话认识" in text
    assert "核心看点" in text


def test_plain_text_export_removes_common_inline_markdown() -> None:
    request = TaskRequest(scene="非遗问答", raw_request="介绍粤剧")
    answer = "## 推荐\n**粤剧**，参考 [官方资料](https://example.com)，关键词 `水袖`。"

    text = build_plain_text(request, answer)

    assert "**" not in text
    assert "[官方资料](https://example.com)" not in text
    assert "`水袖`" not in text
    assert "粤剧" in text
    assert "官方资料" in text
    assert "水袖" in text


def test_docx_export_converts_standard_markdown_table() -> None:
    request = TaskRequest(scene="游客路线", raw_request="安排广州非遗一日游")
    answer = (
        "## 行程安排\n"
        "| 时间 | 内容 |\n"
        "| --- | --- |\n"
        "| 09:00 | 粤剧体验 |\n"
        "| 14:00 | 广绣体验 |\n"
    )

    data = build_docx(request, answer)
    document = Document(BytesIO(data))

    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "时间"
    assert table.cell(0, 1).text == "内容"
    assert table.cell(1, 0).text == "09:00"
    assert table.cell(1, 1).text == "粤剧体验"
    assert table.cell(2, 0).text == "14:00"
    assert table.cell(2, 1).text == "广绣体验"
    assert not any("| --- |" in paragraph.text for paragraph in document.paragraphs)


def test_docx_export_strips_inline_markdown_inside_tables() -> None:
    request = TaskRequest(scene="学生研学", raw_request="设计研学任务")
    answer = (
        "| 任务 | 要求 |\n"
        "| :--- | ---: |\n"
        "| **观察** | 记录 `纹样` |\n"
    )

    document = Document(BytesIO(build_docx(request, answer)))

    assert document.tables[0].cell(1, 0).text == "观察"
    assert document.tables[0].cell(1, 1).text == "记录 纹样"
